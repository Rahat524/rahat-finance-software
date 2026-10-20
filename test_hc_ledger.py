import io,json,sqlite3,tempfile,unittest
from pathlib import Path
from functools import wraps
from flask import Flask,g,jsonify,session
from openpyxl import Workbook,load_workbook
from hc_ledger import register_hc_ledger,calculate,number,import_excel,export_excel,ROOT

def fixture():
    w=Workbook();s=w.active;s.title="Ledger HC"
    s["A5"]="Doc. Date";s["F5"]="Receipts";s["B2"]=5700000;s["B3"]="=1000+50"
    for n,v in enumerate(["2026-10-01","2026-10-01","Opening balance","Receipt","LAHORE BAHRIA - POS","",""],1):s.cell(6,n,v)
    for n,v in enumerate(["2026-10-02","2026-10-02","1st Collection","Receipt","LAHORE BAHRIA - POS",2000,""],1):s.cell(7,n,v)
    for n,v in enumerate(["2026-10-02","2026-10-02","Deposit","Payment","DEPOSIT - BANK ALFALAH 40","","=100+20"],1):s.cell(8,n,v)
    w.create_sheet("Deposit")["A1"]="Preserved deposit tab"
    out=io.BytesIO();w.save(out);return out.getvalue()

class HCLedgerTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=ROOT);self.dbpath=Path(self.temp.name)/"test.db"
        self.app=Flask(__name__);self.app.secret_key="test-session-key"
        def db():
            c=sqlite3.connect(self.dbpath);c.row_factory=sqlite3.Row;return c
        def require_permission(permission):
            def deco(fn):
                @wraps(fn)
                def view(*a,**kw):
                    if permission not in session.get("permissions",[]):return jsonify(error="Permission denied"),403
                    return fn(*a,**kw)
                return view
            return deco
        register_hc_ledger(self.app,db,require_permission,lambda *a:None)
        self.client=self.app.test_client()
        with self.client.session_transaction() as s:s["permissions"]=["cash_view","cash_import","export_data"]
        self.raw=fixture()
    def tearDown(self):self.temp.cleanup()
    def state(self):return self.client.get("/api/hc-workbook").get_json()
    def upload(self,revision=0,csrf=None):
        token=csrf if csrf is not None else self.state()["csrf"]
        return self.client.post("/api/hc-workbook/import",data={"revision":str(revision),"file":(io.BytesIO(self.raw),"HC.xlsx")},headers={"X-HC-CSRF":token})
    def test_import_calculate_and_preserve_tabs(self):
        self.assertEqual(self.upload().status_code,200)
        state=self.state();self.assertEqual(len(state["rows"]),3)
        self.assertEqual(state["summary"]["opening"],1050)
        self.assertEqual(state["summary"]["closing"],2930)
        self.assertEqual(state["rows"][2][17],2000)
        self.assertEqual(state["last_rows"],[0,2])
        response=self.client.get("/api/hc-workbook/export")
        book=load_workbook(io.BytesIO(response.data));self.assertEqual(book["Deposit"]["A1"].value,"Preserved deposit tab")
        self.assertEqual(book["Ledger HC"]["G8"].value,120)
        self.assertTrue(book["Ledger HC"]["H8"].value.startswith("=ROUND(SUM(H7)"))
        self.assertEqual(book["Ledger HC"]["B3"].value,1050)
    def test_export_balance_ignores_empty_amount_text(self):
        self.upload()
        book=load_workbook(io.BytesIO(self.client.get('/api/hc-workbook/export').data))
        sheet=book['Ledger HC']
        # Excel subtraction of empty inline strings produces #VALUE!.
        # SUM handles both blank cells and numeric amounts without masking errors.
        self.assertIn('SUM(F7)-SUM(G7)',sheet['H7'].value)
        self.assertIn('SUM(F8)-SUM(G8)',sheet['H8'].value)
        self.assertIn('SUM(H7)',sheet['H8'].value)
        self.assertIn('ROUND(H8,2)',sheet['I8'].value)
        self.assertEqual(sheet['F7'].value,2000)
        self.assertEqual(sheet['G8'].value,120)

    def test_atomic_paste_and_calculated_cell_protection(self):
        self.upload();s=self.state()
        def paste(rows,targets,cols,revision=None):
            return self.client.patch('/api/hc-workbook',json={'action':'paste','rows':rows,'targets':targets,'cols':cols,'revision':self.state()['revision'] if revision is None else revision},headers={'X-HC-CSRF':s['csrf']})
        self.assertEqual(paste([['2,991,000.00','',999999]],[1],[5,6,7]).status_code,200)
        self.assertEqual(self.state()['rows'][1][5],2991000)
        self.assertEqual(self.state()['rows'][1][7],2992050)
        self.assertEqual(paste([['','-120.00']],[2],[5,6]).status_code,200)
        self.assertEqual(self.state()['rows'][2][6],120)
        revision=self.state()['revision']
        self.assertEqual(paste([[25],['invalid']],[1,2],[5]).status_code,400)
        self.assertEqual(self.state()['revision'],revision)
        self.assertEqual(self.state()['rows'][1][5],2991000)
        self.assertEqual(paste([[30]],[1],[5],revision-1).status_code,409)
        self.assertEqual(paste([[1]],[0],[5]).status_code,400)

    def test_paste_at_end_appends_without_replacing_old_rows(self):
        self.upload();s=self.state();before=s['rows'];start=len(before)
        rows=[['10-Aug-26','11-Aug-26','New collection','Receipt','',2991000,''],['11-Aug-26','12-Aug-26','Next collection','Receipt','',100,'']]
        response=self.client.patch('/api/hc-workbook',json={'action':'paste','revision':s['revision'],'rows':rows,'targets':[start,start+1],'cols':list(range(7))},headers={'X-HC-CSRF':s['csrf']})
        self.assertEqual(response.status_code,200,response.get_json())
        after=self.state()['rows'];self.assertEqual(after[:start],before)
        self.assertEqual(len(after),start+2);self.assertEqual(after[start][5],2991000)
        s=self.state()
        response=self.client.patch('/api/hc-workbook',json={'action':'paste','revision':s['revision'],'rows':[rows[1]],'targets':[len(after)],'cols':list(range(7))},headers={'X-HC-CSRF':s['csrf']})
        self.assertEqual(response.status_code,200);self.assertEqual(len(self.state()['rows']),start+3)

    def test_selected_delete_is_atomic_and_recoverable(self):
        self.upload();s=self.state();headers={'X-HC-CSRF':s['csrf']}
        def delete(rows,revision):return self.client.patch('/api/hc-workbook',json={'action':'delete_many','rows':rows,'revision':revision},headers=headers)
        self.assertEqual(delete([1,0],s['revision']).status_code,400)
        self.assertEqual(self.state()['rows'],s['rows'])
        self.assertEqual(delete([1,2],s['revision']-1).status_code,409)
        self.assertEqual(delete([1,2],s['revision']).status_code,200)
        after=self.state();self.assertEqual(after['rows'][0],s['rows'][0])
        self.assertEqual(after['rows'][1][:7],['']*7)
        self.assertEqual(after['summary']['receipts'],0);self.assertEqual(after['summary']['payments'],0)
        response=self.client.patch('/api/hc-workbook',json={'action':'restore','restore_revision':s['revision'],'revision':after['revision']},headers=headers)
        self.assertEqual(response.status_code,200);self.assertEqual(self.state()['rows'],s['rows'])

    def test_edit_updates_balance_and_rejects_stale_save(self):
        self.upload();state=self.state()
        payload={"action":"cell","revision":1,"row":1,"col":5,"value":"2,500"}
        response=self.client.patch("/api/hc-workbook",json=payload,headers={"X-HC-CSRF":state["csrf"]})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.state()["summary"]["closing"],3430)
        stale=self.client.patch("/api/hc-workbook",json=payload,headers={"X-HC-CSRF":state["csrf"]})
        self.assertEqual(stale.status_code,409)
        self.assertEqual(self.state()["summary"]["closing"],3430)
    def test_permissions_csrf_and_calculated_cells(self):
        self.assertEqual(self.upload(csrf="wrong-token").status_code,403)
        self.upload();state=self.state()
        response=self.client.patch("/api/hc-workbook",json={"action":"cell","revision":1,"row":1,"col":7,"value":0},headers={"X-HC-CSRF":state["csrf"]})
        self.assertEqual(response.status_code,400)
        with self.client.session_transaction() as s:s["permissions"]=["cash_view"]
        self.assertEqual(self.upload(revision=1).status_code,403)
        self.assertEqual(self.client.get("/api/hc-workbook/export").status_code,403)
    def test_last_date_non_contiguous_and_excluded_amounts(self):
        doc=import_excel(self.raw)
        doc["rows"].append(["2026-10-03","2026-10-02","Theft","Receipt","THEFT CASH.",50,"","","","","","","","","","","","","",""])
        x=calculate(doc)
        self.assertEqual(x["summary"]["running_balance"],2930)
        self.assertEqual(x["summary"]["closing"],2980)
        self.assertEqual(x["last_rows"],[0,3])
    def test_expression_safety_and_invalid_import(self):
        with self.assertRaises(ValueError):number("__import__('os').system('whoami')")
        with self.assertRaises(ValueError):number("1/0")
        with self.assertRaises(ValueError):number(float("nan"))
        self.assertEqual(number("=20+5*3"),35)
        self.raw=b"not an excel file";self.assertEqual(self.upload().status_code,400)


class HostedStartupTests(unittest.TestCase):
    def test_gunicorn_import_initializes_missing_schema_and_preserves_database(self):
        import os,subprocess,sys
        with tempfile.TemporaryDirectory(dir=ROOT) as temp:
            script = """
import os,sys,sqlite3
from pathlib import Path
os.environ['LOCALAPPDATA']=sys.argv[1]
sys.path.insert(0,sys.argv[2])
import app
client=app.app.test_client()
response=client.post('/api/login',json={'username':'__deployment_check__','password':''})
assert response.status_code==200 and response.get_json()['ok'] is False
c=app.db()
c.execute('CREATE TABLE startup_sentinel(value TEXT)')
c.execute("INSERT INTO startup_sentinel VALUES('preserve')")
c.commit();c.close()
app.initialize_hosted_database()
c=app.db();assert c.execute('SELECT value FROM startup_sentinel').fetchone()[0]=='preserve';c.close()
"""
            env=os.environ.copy()
            env["PYTHONPATH"]=os.pathsep.join(sys.path)
            result=subprocess.run([sys.executable,"-c",script,temp,str(ROOT)],env=env,capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)

if __name__=="__main__":unittest.main()
