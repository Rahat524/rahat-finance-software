"""Linked input sheets for the HC workbook. All changes share its revision/history."""
import copy, io, re
from datetime import date, datetime
from collections import Counter
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import FormulaRule

NOTES=[5000,1000,500,100,75,50,20,10,5,2,1]
BANK_NOTES=[5000,1000,500,100,50,20,10]
SALARY_TYPES=['SALARIES PAYABLE','ADVANCE SALARY PAYABLE','FINAL SETTLEMENT PAYABLE','UNCLAIMED SALARIES PAYA','UNCLAIMED ADVANCE SALARIE','UNCLAIMED FINAL SETTLEMEN']
SALARY_HEADERS=['EMP_ID','Employee Name','Designation','Branch','Salary Type','CNIC','Mobile#','Net Salary','Email Rec. Date','Status (Paid/Un-Paid)','Paid Date','Paid Voucher No','UnPaid Voucher','Auto Payment Narration','Character Count','Duplicate Check']

def defaults():
 return {'deno':{'date':date.today().isoformat(),'counts':['']*11,'fake':['']*3,'pending':['']*3},'deposit':{},'salary':[]}

def extract(book,document,number,date_text):
 result=defaults()
 def val(s,addr):
  x=s[addr].value
  return number(x) if x not in (None,'') else ''
 if 'Deno HC' in book:
  s=book['Deno HC'];d=result['deno']
  for key,addrs in [('counts',['D'+str(r) for r in range(7,18)]),('fake',['D'+str(r) for r in range(21,24)]),('pending',['E'+str(r) for r in range(26,29)])]:d[key]=[val(s,a) for a in addrs]
  if isinstance(s['C4'].value,(date,datetime)):d['date']=date_text(s['C4'].value)
 if 'Deposit' in book:
  s=book['Deposit'];idx=[i for i,r in enumerate(document['rows']) if r[4].upper() in {'DEPOSIT - BANK ALFALAH 40','DEPOSIT - BANK ALFALAH 940'}]
  for k,i in enumerate(idx):result['deposit'][str(i)]=[val(s,chr(67+c)+str(k+3)) for c in range(7)]
 if 'Salary Details' in book:
  for cells in book['Salary Details'].iter_rows(min_row=3,max_col=13,values_only=True):
   if not any(x is not None for x in cells):continue
   row=[]
   for c,x in enumerate(cells):
    if c in (8,10):row.append(date_text(x) if x else '')
    elif c==7:row.append(number(x) if x is not None else '')
    else:row.append(str(x) if x is not None else '')
   result['salary'].append(row)
 return result

def salary_result(rows,number):
 out=[];seen=set();paid=unpaid=0
 for source in rows:
  r=list(source);tp=r[4].upper();unclaimed=tp.startswith(('UNCLAIM','UNCLAMED'));r[9]='Un-paid' if unclaimed else 'Paid'
  amount=number(r[7]);unpaid+=amount if unclaimed else 0;paid+=0 if unclaimed else amount
  period='';label='Final Set' if 'FINAL' in tp else 'Adv Pay' if 'ADVANCE' in tp else 'Sal Pay'
  if unclaimed:label='Uncl Final' if 'FINAL' in tp else 'Uncl Adv' if 'ADVANCE' in tp else 'Uncl Sal'
  if r[8]:
   dt=date.fromisoformat(r[8]);year,month=dt.year,dt.month
   if 'FINAL' not in tp and 'ADVANCE' not in tp and dt.day<=15:month-=1
   if month==0:year-=1;month=12
   period=f'{year:04d}{month:02d}'
   if 'FINAL' not in tp:period_label=' FTMO '+date(year,month,1).strftime('%b-%y')
   else:period_label=''
  else:period_label=''
  branch='JBG' if r[3].upper() in ('LWH','JBG') else 'L-4' if r[3].upper() in ('L-4','LAHORE-4') else r[3]
  prefix=('Unpaid ' if unclaimed else 'Cash Pd ')+r[0]+' ';suffix=' '+label+period_label+' '+branch
  narration=prefix+r[1][:max(0,50-len(prefix)-len(suffix))]+suffix if r[0] and r[1] else ''
  key=(r[0],tp,period,amount) if r[0] and tp and period and r[7]!='' else None
  status='DUPLICATE' if key in seen else 'OK' if key else ''
  if key:seen.add(key)
  out.append(r+[narration,len(narration) if narration else '',status])
 return {'rows':out,'summary':{'Paid Amount Total':round(paid,2),'Unclaimed Amount Total':round(unpaid,2),'Total Amount':round(paid+unpaid,2)}}

def view(document,name,computed,number):
 aux=document.get('sheets') or defaults()
 if name=='Deno HC':
  d=aux['deno'];rows=[]
  for group,notes in [('counts',NOTES),('fake',[5000,1000,500])]:
   for i,note in enumerate(notes):rows.append([group,note,d[group][i],round(note*number(d[group][i]),2)])
  cash=sum(r[3] for r in rows if r[0]=='counts');fake=sum(r[3] for r in rows if r[0]=='fake');pending=sum(number(x) for x in d['pending']);grand=round(cash+fake+pending,2)
  return {'headers':['Section','Note / Particulars','Count / Amount','Total Amount'],'rows':rows+[['pending',label,d['pending'][i],number(d['pending'][i])] for i,label in enumerate(['Return Counter','Cash Counter Opening','IOU'])],'editable':[2],'numeric':[2,3],'date':d['date'],'summary':{'Cash Total':cash,'Fake Notes':fake,'Pending Transactions':pending,'Grand Total':grand,'SAP Balance':computed['summary']['closing'],'Variance':round(grand-computed['summary']['closing'],2)}}
 if name=='Deposit':
  rows=[];ids=[]
  for i,r in enumerate(document['rows']):
   if r[4].upper() not in {'DEPOSIT - BANK ALFALAH 40','DEPOSIT - BANK ALFALAH 940'}:continue
   counts=aux['deposit'].get(str(i),['']*7);total=round(sum(number(v)*n for v,n in zip(counts,BANK_NOTES)),2);amount=number(r[6]);diff=round(total-amount,2)
   remarks='Match Deposit' if diff==0 else ('Excess Deposit Rs. ' if diff<0 else 'Short Deposit Rs. ')+f'{abs(diff):,.2f}'
   ids.append(i);rows.append([len(rows)+1,r[1]]+counts+[total,amount,diff,remarks])
  return {'headers':['Sr No','Posting Date']+list(map(str,BANK_NOTES))+['Total Amount','Deposit Amount','Difference','Remarks'],'rows':rows,'ids':ids,'editable':list(range(2,9)),'numeric':[0]+list(range(2,12)),'dates':[1],'summary':{'Total Counted':round(sum(r[9] for r in rows),2),'Total Deposited':round(sum(r[10] for r in rows),2),'Difference':round(sum(r[11] for r in rows),2)}}
 if name=='Salary Details':return {'headers':SALARY_HEADERS,'editable':[0,1,2,3,4,5,6,7,8,10,11,12],'numeric':[7,14],'dates':[8,10],'salary_types':SALARY_TYPES,**salary_result(aux['salary'],number)}
 if name=='DataSheet':
  rows=[[n,t] for t in ['Payment','Receipt'] for n in document['offsets'].get(t,[])]
  return {'headers':['Offset Account (Btrans.)','Payment / Receipt'],'rows':rows,'editable':[0,1],'numeric':[],'summary':{'Payment Accounts':len(document['offsets'].get('Payment',[])),'Receipt Accounts':len(document['offsets'].get('Receipt',[]))}}
 raise ValueError('Unknown workbook sheet')

def mutate(document,payload,number,date_text):
 name=payload['sheet'];aux=document.setdefault('sheets',defaults());action=payload.get('sheet_action','cell')
 if name=='DataSheet':
  if action=='append':
   text=str(payload.get('name','')).strip();kind=payload.get('kind')
   if kind not in ('Payment','Receipt') or not text or len(text)>150:raise ValueError('Enter an account name and Payment or Receipt')
   if text not in document['offsets'][kind]:document['offsets'][kind].append(text)
   return
  old=view(document,name,{},number)['rows'];i=int(payload['row']);col=int(payload['col'])
  if not 0<=i<len(old) or col not in (0,1):raise ValueError('Invalid account cell')
  current,kind=old[i];new=str(payload.get('value','')).strip()
  if col==1:
   if new not in ('Payment','Receipt'):raise ValueError('Select Payment or Receipt')
   document['offsets'][kind].remove(current)
   if current not in document['offsets'][new]:document['offsets'][new].append(current)
  else:
   if not new or len(new)>150:raise ValueError('Enter an account name')
   document['offsets'][kind][document['offsets'][kind].index(current)]=new
  return
 if name=='Salary Details':
  rows=aux['salary']
  if action=='append':
   if len(rows)>=10000:raise ValueError('Maximum 10,000 salary entries')
   row=payload['values']
   if not isinstance(row,list) or len(row)!=13:raise ValueError('Salary entry must have 13 input columns')
   index=len(rows)
  else:
   index=int(payload['row']);col=int(payload['col'])
   if not 0<=index<len(rows) or not 0<=col<13:raise ValueError('This cell is calculated or does not exist')
   row=rows[index].copy();row[col]=payload.get('value','')
  row=[str(v or '').strip() if c!=7 else (number(v) if v not in ('',None) else '') for c,v in enumerate(row)]
  for c in (8,10):row[c]=date_text(row[c])
  if row[4] not in SALARY_TYPES and row[4] not in [r[4] for r in rows]:raise ValueError('Select a Salary Type')
  if number(row[7])<0:raise ValueError('Salary amount cannot be negative')
  row[9]='Un-paid' if row[4].startswith(('UNCLAIM','UNCLAMED')) else 'Paid'
  if any(len(str(v))>2000 for v in row):raise ValueError('Cell text is too long')
  if index==len(rows):rows.append(row)
  else:rows[index]=row
  return
 if name=='Deno HC' and action=='date':
  dt=date_text(payload['value'])
  if not dt:raise ValueError('Select a denomination date')
  aux['deno']['date']=dt;return
 value=number(payload.get('value','')) if payload.get('value','')!='' else ''
 if value!='' and value<0:raise ValueError('Counts and amounts cannot be negative')
 index=int(payload['row']);col=int(payload['col'])
 if name=='Deno HC':
  if col!=2 or not 0<=index<17:raise ValueError('Only count and pending amount cells are editable')
  group,idx=('counts',index) if index<11 else ('fake',index-11) if index<14 else ('pending',index-14)
  if group!='pending' and value!='' and int(value)!=value:raise ValueError('Note counts must be whole numbers')
  aux['deno'][group][idx]=value
 elif name=='Deposit':
  if not 2<=col<=8 or not 0<=index<len(document['rows']):raise ValueError('Invalid deposit cell')
  if document['rows'][index][4].upper() not in {'DEPOSIT - BANK ALFALAH 40','DEPOSIT - BANK ALFALAH 940'}:raise ValueError('This ledger row is not a bank deposit')
  if value!='' and int(value)!=value:raise ValueError('Note counts must be whole numbers')
  aux['deposit'].setdefault(str(index),['']*7)[col-2]=value
 else:raise ValueError('Unknown workbook sheet')

def export_aux(book,document,number):
 aux=document.get('sheets')
 if not aux:return
 if 'Deno HC' not in book:
  ds=book.create_sheet('Deno HC');ds['B2']='DAILY CASH DENOMINATION (Head Cashier)';ds['B3']='Imtiaz Super Markets - Lahore Bahria Town L-4'
  for r,n in list(enumerate(NOTES,7))+list(enumerate([5000,1000,500],21)):
   ds.cell(r,2,n);ds.cell(r,3,'x');ds.cell(r,5,f'=B{r}*D{r}')
  for a,f in {'E18':'=SUM(E7:E17)','E24':'=SUM(E21:E23)','E29':'=SUM(E26:E28)','E31':'=E18+E24+E29','E32':"='Ledger HC'!B4",'E33':'=E31-E32'}.items():ds[a]=f
  for r,l in [(26,'Return Counter'),(27,'Cash Counter Opening'),(28,'IOU'),(31,'Grand Total'),(32,'SAP Balance'),(33,'Variance')]:ds.cell(r,3,l)
 if 'Deno HC' in book:
  s=book['Deno HC'];d=aux['deno'];s['C4']=datetime.fromisoformat(d['date']);s['C4'].number_format='dd-mmm-yy'
  for group,addrs in [('counts',['D'+str(r) for r in range(7,18)]),('fake',['D'+str(r) for r in range(21,24)]),('pending',['E'+str(r) for r in range(26,29)])]:
   for addr,v in zip(addrs,d[group]):s[addr]=v if v!='' else None
 if 'Deposit' not in book:
  ds=book.create_sheet('Deposit');ds['A1']='Deposit Details - Bank Alfalah'
  for c,v in enumerate(['Sr No','Posting Date']+list(map(str,BANK_NOTES))+['Total Amount','Deposit Amount','Difference','Remarks'],1):ds.cell(2,c,v)
 if 'Deposit' in book:
  s=book['Deposit'];bank=[i for i,r in enumerate(document['rows']) if r[4].upper() in {'DEPOSIT - BANK ALFALAH 40','DEPOSIT - BANK ALFALAH 940'}]
  for seq,i in enumerate(bank,3):
   r=seq;ledger=i+6
   s.cell(r,1,seq-2);s.cell(r,2,f"='Ledger HC'!B{ledger}")
   s.cell(r,10,f'=SUMPRODUCT(C{r}:I{r},{{5000,1000,500,100,50,20,10}})')
   s.cell(r,11,f"='Ledger HC'!G{ledger}");s.cell(r,12,f'=ROUND(J{r}-K{r},2)')
   s.cell(r,13,f'=IF(L{r}=0,"Match Deposit",IF(L{r}<0,"Excess Deposit Rs. ","Short Deposit Rs. ")&TEXT(ABS(L{r}),"#,##0.00"))')
   for col,v in enumerate(aux['deposit'].get(str(i),['']*7),3):s.cell(seq,col,v if v!='' else None).value=v if v!='' else None
 if 'Salary Details' not in book:
  ds=book.create_sheet('Salary Details')
  for c,v in enumerate(SALARY_HEADERS+['Duplicate Key (Helper)'],1):ds.cell(2,c,v)
 if 'Salary Details' in book:
  s=book['Salary Details'];result=salary_result(aux['salary'],number)
  end=max(3,len(result['rows'])+2)
  s['B1']='Paid Amount Total';s['D1']=f'=SUMIF(J3:J{end},"Paid",H3:H{end})'
  s['G1']='Total Amount';s['H1']=f'=SUM(H3:H{end})'
  s['J1']='Unclaimed Amount Total';s['L1']=f'=SUMIF(J3:J{end},"Un-paid",H3:H{end})'
  s.column_dimensions['Q'].hidden=True
  s.conditional_formatting.add(f'A3:P{max(3000,end)}',FormulaRule(formula=['$P3="DUPLICATE"'],fill=PatternFill('solid',fgColor='FFE0DC'),font=Font(color='9B2017')))
  narration_template=s['N3'].value if s['N3'].data_type=='f' else None
  for i,r in enumerate(result['rows'],3):
   native={c:s.cell(i,c).value for c in (14,15) if s.cell(i,c).data_type=='f'}
   for col,v in enumerate(r[:13],1):
    if col in (9,11) and v:v=datetime.fromisoformat(v)
    c=s.cell(i,col);c.value=v if v!='' else None
    if isinstance(v,str):c.data_type='s'
    if col in (9,11):c.number_format='dd-mmm-yy'
    if col==8:c.number_format='#,##0.00'
   s.cell(i,14,native.get(14) or (Translator(narration_template,origin='N3').translate_formula('N'+str(i)) if narration_template else r[13]))
   s.cell(i,15,native.get(15) or f'=IF(N{i}="","",LEN(N{i}))')
   s.cell(i,17,f'=IF(OR(A{i}="",E{i}="",I{i}="",H{i}=""),"",A{i}&"|"&UPPER(TRIM(E{i}))&"|"&TEXT(IF(OR(ISNUMBER(SEARCH("FINAL",E{i})),ISNUMBER(SEARCH("ADVANCE",E{i})),DAY(I{i})>15),I{i},EDATE(I{i},-1)),"yyyymm")&"|"&TEXT(H{i},"0.00"))')
   s.cell(i,16,f'=IF(Q{i}="","",IF(COUNTIF($Q$3:Q{i},Q{i})>1,"DUPLICATE","OK"))')
 if 'DataSheet' not in book:book.create_sheet('DataSheet')
 if 'DataSheet' in book:
  s=book['DataSheet']
  for j,kind in [(1,'Payment'),(6,'Receipt')]:
   names=document['offsets'][kind]
   for i in range(2,max(s.max_row,len(names)+1)+1):
    s.cell(i,j).value=names[i-2] if i-2<len(names) else None;s.cell(i,j+3).value=kind if i-2<len(names) else None
    if s.cell(i,j).value is not None:s.cell(i,j).data_type='s'
