import io, unittest, copy
from openpyxl import load_workbook
from hc_ledger import calculate, import_excel, export_excel, number, date_text, ranked_offsets
from hc_sheets import defaults, mutate, view, salary_result
from test_hc_ledger import HCLedgerTests

class CompletionTests(HCLedgerTests):
 def test_linked_inputs_revision_and_export(self):
  self.upload();s=self.state();token=s['csrf']
  def patch(**p):return self.client.patch('/api/hc-workbook',json={'action':'sheet','revision':self.state()['revision'],**p},headers={'X-HC-CSRF':token})
  self.assertEqual(patch(sheet='Deno HC',row=0,col=2,value='2+3').status_code,200)
  d=self.client.get('/api/hc-workbook/sheet?name=Deno%20HC').get_json()
  self.assertEqual(d['summary']['Cash Total'],25000)
  self.assertEqual(d['summary']['SAP Balance'],2930)
  self.assertEqual(patch(sheet='Deno HC',row=0,col=3,value=5).status_code,400)
  self.assertEqual(patch(sheet='Deno HC',row=0,col=2,value='0.5').status_code,400)
  self.assertEqual(patch(sheet='Deposit',row=2,col=2,value=1).status_code,200)
  d=self.client.get('/api/hc-workbook/sheet?name=Deposit').get_json()
  self.assertEqual(d['rows'][0][9:12],[5000,120,4880])
  self.assertEqual(patch(sheet='DataSheet',sheet_action='append',name='NEW TEST ACCOUNT',kind='Payment').status_code,200)
  self.assertIn('NEW TEST ACCOUNT',self.state()['offsets']['Payment'])
  self.assertNotIn('NEW TEST ACCOUNT',self.state()['offsets']['Receipt'])
  book=load_workbook(io.BytesIO(self.client.get('/api/hc-workbook/export').data))
  self.assertEqual(book['Deno HC']['D7'].value,5)
  self.assertEqual(book['Deposit']['C3'].value,1)
  self.assertEqual(book['Deposit']['K3'].value,"='Ledger HC'!G8")
  self.assertTrue(book['Ledger HC']['G8'].number_format.startswith('"-"'))
  self.assertIn('ROUND',book['Ledger HC']['H8'].value)
  self.assertIn('ADVANCE SALARY PAYABLE',book['Ledger HC']['R8'].value)
 def test_new_endpoints_permissions_and_stale(self):
  self.upload();s=self.state()
  for name in ['Deno HC','Deposit','Salary Details','DataSheet']:
   self.assertEqual(self.client.get('/api/hc-workbook/sheet',query_string={'name':name}).status_code,200)
  self.assertEqual(self.client.get('/api/hc-workbook/sheet?name=unknown').status_code,400)
  p={'action':'sheet','sheet':'Deno HC','row':0,'col':2,'value':2,'revision':1}
  self.assertEqual(self.client.patch('/api/hc-workbook',json=p,headers={'X-HC-CSRF':'bad'}).status_code,403)
  self.assertEqual(self.client.patch('/api/hc-workbook',json=p,headers={'X-HC-CSRF':s['csrf']}).status_code,200)
  self.assertEqual(self.client.patch('/api/hc-workbook',json=p,headers={'X-HC-CSRF':s['csrf']}).status_code,409)
  with self.client.session_transaction() as x:x['permissions']=[]
  self.assertEqual(self.client.get('/api/hc-workbook/sheet?name=Deno%20HC').status_code,403)
 def test_salary_and_unclaimed_deposit_not_sale(self):
  doc=import_excel(self.raw)
  def row(offset,f='',g='',text='Test',day='2026-10-02'):
   return [day,day,text,'Receipt' if f else 'Payment',offset,f,g]+['']*13
  doc['rows'] += [row('ADVANCE SALARY PAYABLE',g=100),row('SALARIES PAYABLE',g=200),row('FINAL SETTLEMENT PAYABLE',g=300),row('SALARIES PAYABLE',g=999,day='2026-10-03'),row('UNCLAIMED SALARIES PAYABL',f=250),row('DEPOSIT - BANK ALFALAH 40',g=100,text='Cash Depot (940) Unclimd Salry test')]
  result=calculate(doc)
  self.assertEqual(result['rows'][2][17],1400)
  self.assertEqual(result['rows'][-1][17],'')
  self.assertEqual(result['summary']['unclaimed_received'],250)
  self.assertEqual(result['summary']['unclaimed_deposited'],100)
  self.assertEqual(result['summary']['unclaimed_remaining'],150)
  self.assertEqual(result['summary']['remaining'],150)
  doc['rows'][-1][4]='DEPOSIT - BANK ALFALAH 66'
  self.assertEqual(calculate(doc)['summary']['remaining'],150)
 def test_salary_types_duplicate_boundaries_and_export(self):
  self.upload();token=self.state()['csrf']
  def add(tp):
   vals=['123','Test Employee','','L-4',tp,'','',1000,'2026-10-20','','2026-10-20','','']
   return self.client.patch('/api/hc-workbook',json={'action':'sheet','sheet':'Salary Details','sheet_action':'append','values':vals,'revision':self.state()['revision']},headers={'X-HC-CSRF':token})
  for tp in ['FINAL SETTLEMENT PAYABLE','UNCLAIMED FINAL SETTLEMEN','FINAL SETTLEMENT PAYABLE','UNCLAIMED FINAL SETTLEMEN']:
   self.assertEqual(add(tp).status_code,200)
  data=self.client.get('/api/hc-workbook/sheet?name=Salary%20Details').get_json()
  self.assertEqual([r[15] for r in data['rows']],['OK','OK','DUPLICATE','DUPLICATE'])
  self.assertEqual([r[9] for r in data['rows']],['Paid','Un-paid','Paid','Un-paid'])
  self.assertTrue(data['rows'][1][13].startswith('Unpaid '))
  self.assertEqual(data['summary']['Paid Amount Total'],2000)
  self.assertEqual(data['summary']['Unclaimed Amount Total'],2000)
  b=load_workbook(io.BytesIO(self.client.get('/api/hc-workbook/export').data))
  self.assertIn('TEXT',b['Salary Details']['Q3'].value)
 def test_delete_recoverable_and_rank_offsets(self):
  self.upload();s=self.state()
  self.assertEqual(self.client.patch('/api/hc-workbook',json={'action':'delete','row':0,'revision':1},headers={'X-HC-CSRF':s['csrf']}).status_code,400)
  self.assertEqual(self.client.patch('/api/hc-workbook',json={'action':'delete','row':2,'revision':1},headers={'X-HC-CSRF':s['csrf']}).status_code,200)
  self.assertEqual(self.state()['summary']['payments'],0)
  doc=import_excel(self.raw);doc['offsets']['Payment']=['Unused','DEPOSIT - BANK ALFALAH 40']
  self.assertEqual(ranked_offsets(doc)['Payment'][0],'DEPOSIT - BANK ALFALAH 40')
 def test_history_restore_keeps_current_revision(self):
  self.upload();s=self.state();headers={'X-HC-CSRF':s['csrf']}
  r=self.client.patch('/api/hc-workbook',json={'action':'cell','row':1,'col':5,'value':2500,'revision':1},headers=headers)
  self.assertEqual(r.status_code,200)
  self.assertEqual(self.state()['summary']['closing'],3430)
  history=self.client.get('/api/hc-workbook/history').get_json()['entries']
  self.assertIn(1,[r['revision'] for r in history])
  r=self.client.patch('/api/hc-workbook',json={'action':'restore','restore_revision':1,'revision':2},headers=headers)
  self.assertEqual(r.status_code,200)
  self.assertEqual(self.state()['summary']['closing'],2930)
  self.assertEqual(self.state()['revision'],3)
 def test_export_preserves_original_design_and_other_sheet(self):
  from openpyxl.styles import Font,PatternFill
  b=load_workbook(io.BytesIO(self.raw));s=b['Ledger HC']
  s.column_dimensions['C'].width=71;s.column_dimensions['J'].hidden=False
  s['C5'].font=Font(name='Arial',size=14,bold=True,color='FF00AA')
  s['C5'].fill=PatternFill('solid',fgColor='BBCCDD');s['A7'].number_format='d-mmm-yyyy'
  s.freeze_panes='C6';s.sheet_view.zoomScale=85;s.protection.sheet=True
  b['Deposit']['B1']='Unrelated value retained'
  output=io.BytesIO();b.save(output);self.raw=output.getvalue();self.upload()
  result=load_workbook(io.BytesIO(self.client.get('/api/hc-workbook/export').data));out=result['Ledger HC']
  self.assertEqual(out.column_dimensions['C'].width,71)
  self.assertFalse(out.column_dimensions['J'].hidden)
  self.assertEqual(copy.copy(out['C5'].font),copy.copy(s['C5'].font))
  self.assertEqual(copy.copy(out['C5'].fill),copy.copy(s['C5'].fill))
  self.assertEqual(out['A7'].number_format,'d-mmm-yyyy')
  self.assertEqual(out.freeze_panes,'C6');self.assertEqual(out.sheet_view.zoomScale,85)
  self.assertTrue(out.protection.sheet)
  self.assertEqual(result['Deposit']['B1'].value,'Unrelated value retained')
 def test_fractional_limit_match(self):
  doc=import_excel(self.raw);doc['opening']=5700000.0-0.1-0.2;doc['limit']=5700000;doc['rows'][1][5]=0.3;doc['rows'][2][6]=''
  self.assertEqual(calculate(doc)['rows'][1][8],'Limit Match')

if __name__=='__main__':unittest.main()
