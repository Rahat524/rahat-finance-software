import io,os,sys,tempfile,importlib,unittest,hashlib
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from test_hc_ledger import fixture

class StandaloneTests(unittest.TestCase):
 def test_isolated_login_import_edit_and_excel_export(self):
  root=Path(__file__).resolve().parent;old=root/'corporate_scrap.db'
  before=hashlib.sha256(old.read_bytes()).digest() if old.exists() else None
  with tempfile.TemporaryDirectory(dir=root) as tmp:
   with patch.dict(os.environ,{'HC_DATA_DIR':tmp,'HC_USERNAME':'test-owner','HC_PASSWORD':'test-only-password','APP_SECRET_KEY':'test-only-session-key','RENDER':''}):
    sys.modules.pop('hc_app',None);module=importlib.import_module('hc_app');client=module.app.test_client()
    self.assertEqual(client.get('/').status_code,200)
    self.assertEqual(client.get('/api/hc-workbook').status_code,401)
    for path in ['/hc-ledger','/hc-sheets?name=Deposit']:
     page=client.get(path);self.assertEqual(page.status_code,302);self.assertEqual(page.headers['Location'],'/')
    self.assertFalse(client.post('/api/login',json={'username':'test-owner','password':'wrong'}).get_json()['ok'])
    self.assertTrue(client.post('/api/login',json={'username':'test-owner','password':'test-only-password'}).get_json()['ok'])
    self.assertEqual(client.get('/api/me').get_json()['role_name'],'Super Admin')
    state=client.get('/api/hc-workbook').get_json();headers={'X-HC-CSRF':state['csrf']}
    response=client.post('/api/hc-workbook/import',data={'revision':0,'file':(io.BytesIO(fixture()),'HC.xlsx')},headers=headers)
    self.assertEqual(response.status_code,200)
    self.assertEqual(client.get('/api/hc-workbook').get_json()['summary']['closing'],2930)
    for name in ['Deno HC','Deposit','Salary Details','DataSheet']:
     self.assertEqual(client.get('/api/hc-workbook/sheet',query_string={'name':name}).status_code,200)
    book=load_workbook(io.BytesIO(client.get('/api/hc-workbook/export').data))
    self.assertIn('Ledger HC',book.sheetnames)
    self.assertEqual(book['Ledger HC']['G8'].value,120)
    c=module.db();tables={r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")};c.close()
    self.assertIn('hc_workbook',tables);self.assertNotIn('customer_ledger',tables);self.assertNotIn('vendor_ledger',tables)
    self.assertEqual(Path(module.DB).parent.resolve(),Path(tmp).resolve())
    client.post('/api/logout');self.assertEqual(client.get('/api/hc-workbook').status_code,401)
  if before is not None:self.assertEqual(hashlib.sha256(old.read_bytes()).digest(),before)

if __name__=='__main__':unittest.main()
