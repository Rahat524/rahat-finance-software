"""Excel-layout HC Ledger. Transaction data stays in the application database."""
from __future__ import annotations
import ast, copy, io, json, math, operator, secrets, sqlite3, zipfile
from collections import Counter
from functools import lru_cache
from hc_sheets import extract as extract_sheets, defaults as sheet_defaults, view as sheet_view, mutate as mutate_sheet, export_aux
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from flask import Blueprint, jsonify, request, send_file, send_from_directory, session
from openpyxl import load_workbook, Workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
INPUTS = {0,1,2,3,4,5,6,9,10,11,12,13,14,15,16}
NUMBERS = {5,6,9,10,11,12,13,14,15,16}
SALARY_PAYMENTS = {"SALARIES PAYABLE","ADVANCE SALARY PAYABLE","FINAL SETTLEMENT PAYABLE"}
DEPOSITS = {"DEPOSIT - BANK ALFALAH 40","DEPOSIT - BANK ALFALAH 940"}
OPS = {ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}

def number(value):
    if value in (None, ""): return 0.0
    if isinstance(value, bool): raise ValueError("Enter an amount, not True/False")
    if isinstance(value, (int,float)): result=float(value)
    else:
        text=str(value).strip().replace(",","").lstrip("=")
        if len(text)>4000: raise ValueError("Amount expression is too long")
        def calc(node):
            if isinstance(node,ast.Expression): return calc(node.body)
            if isinstance(node,ast.Constant) and type(node.value) in (int,float): return node.value
            if isinstance(node,ast.BinOp) and type(node.op) in OPS: return OPS[type(node.op)](calc(node.left),calc(node.right))
            if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)): return calc(node.operand)*(1 if isinstance(node.op,ast.UAdd) else -1)
            raise ValueError("Only numbers and + - * / expressions are supported in amount cells")
        try: result=float(calc(ast.parse(text,mode="eval")))
        except (SyntaxError,ZeroDivisionError,RecursionError,OverflowError) as exc: raise ValueError("Invalid amount expression") from exc
    if not math.isfinite(result) or abs(result)>1e15: raise ValueError("Amount is outside the supported range")
    return round(result,2)

def date_text(value):
    if value in (None,""): return ""
    if isinstance(value,(datetime,date)): return value.strftime("%Y-%m-%d")
    text=str(value).strip()
    try: return datetime.fromisoformat(text).strftime("%Y-%m-%d")
    except ValueError:
        for fmt in ("%d-%b-%y","%d-%b-%Y","%d/%m/%Y","%Y-%m-%d"):
            try: return datetime.strptime(text,fmt).strftime("%Y-%m-%d")
            except ValueError: pass
    raise ValueError("Use a valid date (YYYY-MM-DD)")

def sanitize_row(values):
    if not isinstance(values,list) or len(values)!=20: raise ValueError("Each ledger row must have 20 columns")
    row=[""]*20
    for col in INPUTS:
        value=values[col]
        if col in (0,1): row[col]=date_text(value)
        elif col in NUMBERS:
            row[col]=number(value) if value not in (None,"") else ""
            if row[col]!="" and row[col]<0: raise ValueError("Amounts cannot be negative")
        else:
            row[col]=str(value or "").strip()
            if len(row[col])>2000: raise ValueError("Cell text is too long")
    if row[3] not in ("","Payment","Receipt"): raise ValueError("Cash Journal Type must be Payment or Receipt")
    if row[5] and row[6]: raise ValueError("Use either Receipts or Payment in a row")
    return row

def calculate(document):
    rows=copy.deepcopy(document["rows"]); opening=number(document.get("opening",0)); limit=number(document.get("limit",5700000))
    daily={}; last_pos={}; last_date={}
    for index,row in enumerate(rows):
        key=row[0]; offset=row[4].upper(); desc=row[2].upper()
        day=daily.setdefault(key,{"collections":0,"pc":0,"returns":0,"pos":0,"salary":0})
        if offset in SALARY_PAYMENTS: day["salary"]+=number(row[6])
        if "COLLECTION" in desc: day["collections"]+=number(row[5])
        if "TILL CLOSING" in desc: day["collections"]+=number(row[5])
        if "CASH TRNSFR H.C TO P.C" in offset: day["pc"]+=number(row[6])
        if "RETURN POS" in offset: day["returns"]+=number(row[6])
        if offset=="LAHORE BAHRIA - POS" and key:
            day["pos"]+=number(row[5]); last_pos[key]=index
        if row[1]: last_date[row[1]]=index
    balance=opening; receipt=payment=theft=lost=alfalah=unclaimed=unclaimed_deposit=unclaimed_other_bank=0
    for index,row in enumerate(rows):
        f,g=number(row[5]),number(row[6]); receipt+=f;payment+=g;offset=row[4].upper()
        if offset.startswith(("UNCLAIM","UNCLAMED")): unclaimed+=f
        if "DEPOSIT" in offset and "UNCL" in row[2].upper():
            unclaimed_deposit+=g
            if offset!="DEPOSIT - BANK ALFALAH 66":unclaimed_other_bank+=g
        if "THEFT CASH" in offset: theft+=f
        if "LOST & FOUND" in offset: lost+=f
        if offset=="DEPOSIT - BANK ALFALAH 66": alfalah+=g
        # In the source Excel the first data row is an opening-balance row.
        if index!=0 and not any(word in offset for word in ("THEFT CASH","LOST & FOUND","DEPOSIT - BANK ALFALAH 66")): balance=round(balance+f-g,2)
        row[7]=balance
        row[8]="Limit Exceed" if balance>limit else "Limit Match" if round(balance,2)==round(limit,2) else "Below Limit"
        day=daily[row[0]]
        row[17]=round(day["collections"]-day["pc"]-day["returns"]-day["salary"],2) if offset in DEPOSITS and "UNCL" not in row[2].upper() else ""
        row[18]=round(day["pos"],2) if last_pos.get(row[0])==index else ""
        row[19]=len(row[2]) if row[2] else ""
    return {"rows":rows,"last_rows":list(last_date.values()),"summary":{
        "limit":limit,"opening":opening,"receipts":round(receipt,2),"payments":round(payment,2),
        "closing":round(opening+receipt-payment,2),"theft":round(theft,2),"lost":round(lost,2),
        "alfalah66":round(alfalah,2),"remaining":round(theft+lost+unclaimed-alfalah-unclaimed_other_bank,2),
        "unclaimed_received":round(unclaimed,2),"unclaimed_deposited":round(unclaimed_deposit,2),
        "unclaimed_remaining":round(unclaimed-unclaimed_deposit,2),
        "running_balance":balance}}

def import_excel(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist())>80*1024*1024: raise ValueError("Workbook is too large after extraction")
        book=load_workbook(io.BytesIO(raw),data_only=False)
    except (zipfile.BadZipFile,KeyError) as exc: raise ValueError("Select a valid .xlsx workbook") from exc
    if "Ledger HC" not in book.sheetnames: raise ValueError("The workbook must contain Ledger HC")
    sheet=book["Ledger HC"]
    if str(sheet["A5"].value).strip()!="Doc. Date" or str(sheet["F5"].value).strip()!="Receipts": raise ValueError("The Ledger HC column layout does not match")
    if sheet.max_row>10000: raise ValueError("Maximum supported sheet size is 10,000 rows")
    rows=[]
    for cells in sheet.iter_rows(min_row=6,max_col=20):
        if not any(c.value is not None for c in cells[:7]): continue
        try:
            values=[c.value if i in INPUTS else "" for i,c in enumerate(cells)]
            rows.append(sanitize_row(values))
        except ValueError as exc: raise ValueError(f"Excel row {cells[0].row}: {exc}") from exc
    if not rows: raise ValueError("No ledger entries found")
    if number(rows[0][5]) or number(rows[0][6]): raise ValueError("The first ledger row must be the opening-balance row with no receipt/payment")
    offsets={"Payment":[],"Receipt":[]}
    if "DataSheet" in book:
        for cells in book["DataSheet"].iter_rows(min_row=2,values_only=True):
            for name_col in (0,5):
                if len(cells)>name_col+3 and cells[name_col] and cells[name_col+3] in offsets:
                    offsets[cells[name_col+3]].append(str(cells[name_col]))
    document={"opening":number(sheet["B3"].value),"limit":number(sheet["B2"].value),"rows":rows,"offsets":{k:list(dict.fromkeys(v)) for k,v in offsets.items()}}
    document["sheets"]=extract_sheets(book,document,number,date_text)
    book.close()
    return document

def export_excel(document, original, template):
    book=load_workbook(io.BytesIO(original)) if original else Workbook()
    if not original:
        sheet=book.active;sheet.title="Ledger HC"
        sheet.merge_cells("A1:U1");sheet["A1"]=template["title"]
        for address,label in {"A2":"Approved Limit","D2":"Total Receipts","G2":"Total Payments","A3":"Opening Balance","D3":"Theft Cash","G3":"LOST & FOUND.","A4":"Closing Balance","D4":"Deposit - Alfalah 66","G4":"Remaining Amount"}.items(): sheet[address]=label
        for col,header in enumerate(template["headers"],1): sheet.cell(5,col,header)
    else: sheet=book["Ledger HC"]
    rows=document["rows"]; end=max(6,len(rows)+5); old_end=sheet.max_row
    # Retain the original workbook, all other tabs, cell formats and formulas.
    for r in range(6,max(old_end,end)+1):
        for col in INPUTS: sheet.cell(r,col+1).value=None
    sheet["B2"]=document["limit"];sheet["B3"]=document["opening"]
    sheet["E2"]=f"=SUM(F6:F{end})";sheet["H2"]=f"=SUM(G6:G{end})"
    sheet["B4"]="=B3+E2-H2"
    sheet["E3"]=f'=SUMIF(E6:E{end},"*THEFT CASH*",F6:F{end})'
    sheet["H3"]=f'=SUMIF(E6:E{end},"*LOST & FOUND.*",F6:F{end})'
    sheet["E4"]=f'=SUMIF(E6:E{end},"DEPOSIT - BANK ALFALAH 66",G6:G{end})'
    sheet["T2"]=f'=ROUND(SUMIF(E6:E{end},"UNCLAIM*",F6:F{end})+SUMIF(E6:E{end},"UNCLAMED*",F6:F{end}),2)'
    sheet["T3"]=f'=ROUND(SUMIFS(G6:G{end},E6:E{end},"*DEPOSIT*",C6:C{end},"*UNCL*"),2)'
    sheet["T4"]='=ROUND(T2-T3,2)'
    for rr,label in [(2,'Unclaimed Salary Received'),(3,'Unclaimed Salary Deposited'),(4,'Unclaimed Salary Remaining')]:
        sheet.cell(rr,17,label)
        if f'Q{rr}:S{rr}' not in sheet.merged_cells:sheet.merge_cells(f'Q{rr}:S{rr}')
        sheet.cell(rr,17)._style=copy.copy(sheet.cell(rr,7)._style)
        sheet.cell(rr,20)._style=copy.copy(sheet.cell(rr,8)._style)
        sheet.cell(rr,20).number_format='#,##0.00'
    sheet["H4"]=f'=ROUND(E3+H3+T2-E4-SUMIFS(G6:G{end},E6:E{end},"*DEPOSIT*",E6:E{end},"<>DEPOSIT - BANK ALFALAH 66",C6:C{end},"*UNCL*"),2)' 
    computed=calculate(document)
    for index,row in enumerate(rows):
        r=index+6
        for col in INPUTS:
            value=row[col]
            if col in (0,1) and value: value=datetime.strptime(value,"%Y-%m-%d")
            cell=sheet.cell(r,col+1);cell.value=value if value!="" else None
            if isinstance(value,str): cell.data_type="s"
            if col in (0,1) and (not original or r>old_end):cell.number_format="dd-mmm-yy"
        sheet.cell(r,8,"=ROUND($B$3,2)" if index==0 else f'=ROUND(H{r-1}+IF(OR(COUNTIF(E{r},"*THEFT CASH*")>0,COUNTIF(E{r},"*LOST & FOUND*")>0,COUNTIF(E{r},"*DEPOSIT - BANK ALFALAH 66*")>0),0,F{r}-G{r}),2)')
        sheet.cell(r,9,f'=IF(ROUND(H{r},2)>ROUND($B$2,2),"Limit Exceed",IF(ROUND(H{r},2)=ROUND($B$2,2),"Limit Match","Below Limit"))')
        sheet.cell(r,18,f'=IF(AND(OR(E{r}="DEPOSIT - BANK ALFALAH 40",E{r}="DEPOSIT - BANK ALFALAH 940"),COUNTIF(C{r},"*UNCL*")=0),SUMIFS(F6:F{end},A6:A{end},A{r},C6:C{end},"*Collection*")+SUMIFS(F6:F{end},A6:A{end},A{r},C6:C{end},"*Till Closing*")-SUMIFS(G6:G{end},A6:A{end},A{r},E6:E{end},"*CASH TRNSFR H.C TO P.C*")-SUMIFS(G6:G{end},A6:A{end},A{r},E6:E{end},"*RETURN POS*")-SUMIFS(G6:G{end},A6:A{end},A{r},E6:E{end},"SALARIES PAYABLE")-SUMIFS(G6:G{end},A6:A{end},A{r},E6:E{end},"ADVANCE SALARY PAYABLE")-SUMIFS(G6:G{end},A6:A{end},A{r},E6:E{end},"FINAL SETTLEMENT PAYABLE"),"")')
        sheet.cell(r,19,f'=IF(AND(ISNUMBER(A{r}),E{r}="LAHORE BAHRIA - POS"),IF(COUNTIFS(A{r+1}:A{end+1},A{r},E{r+1}:E{end+1},"LAHORE BAHRIA - POS")=0,SUMIFS(F6:F{end},A6:A{end},A{r},E6:E{end},"LAHORE BAHRIA - POS"),""),"")')
        sheet.cell(r,20,f'=IF(C{r}="","",LEN(C{r}))')
        sheet.cell(r,21,f'=IF(ISNUMBER(B{r}),--(COUNTIF(B6:B{end},B{r})=COUNTIF(B6:B{r},B{r})),0)')
        sheet.cell(r,22,('=0' if index==0 else f'=V{r-1}')+f'+--(OR(E{r}="DEPOSIT - BANK ALFALAH 40",E{r}="DEPOSIT - BANK ALFALAH 940"))')
        for c in range(1,21):
            cell=sheet.cell(r,c)
            if not original or r>old_end:
                cell.font=Font(name="Trebuchet MS",size=9)
                cell.fill=PatternFill("solid",fgColor="FFF2CC" if index in computed["last_rows"] else "FFFFFF")
                if c in (6,7,8,18,19): cell.number_format='#,##0.00;[Red](#,##0.00);"-"'
    for rr in range(6,max(end,old_end)+1):sheet.cell(rr,7).number_format='"-"#,##0.00;"-"#,##0.00;0.00'
    export_aux(book,document,number)
    if not original:
        for col in range(1,21):
            letter=get_column_letter(col)
            sheet.column_dimensions[letter].width=template["widths"][col-1]/7
            sheet.column_dimensions[letter].hidden=col-1 in template["hidden"]
            sheet.cell(5,col).fill=PatternFill("solid",fgColor="234668");sheet.cell(5,col).font=Font(name="Trebuchet MS",color="FFFFFF",bold=True,size=9)
            sheet.cell(5,col).alignment=Alignment(wrap_text=True)
        sheet.freeze_panes="F6"
    sheet.auto_filter.ref=f"A5:T{end}"
    from openpyxl.workbook.properties import CalcProperties
    book.calculation=CalcProperties(fullCalcOnLoad=True)
    out=io.BytesIO();book.save(out);out.seek(0);return out

def ranked_offsets(document):
    counts=Counter((r[3],r[4]) for r in document['rows'])
    return {kind:sorted(list(dict.fromkeys(names)),key=lambda name:(-counts[(kind,name)],name.casefold())) for kind,names in document['offsets'].items()}

def register_hc_ledger(app, db, require_permission, audit):
    bp=Blueprint("hc_workbook",__name__)
    template=json.loads((ROOT/"hc_ledger_template.json").read_text(encoding="utf-8"))
    def connect():
        conn=db()
        conn.execute("CREATE TABLE IF NOT EXISTS hc_workbook(id INTEGER PRIMARY KEY CHECK(id=1),revision INTEGER NOT NULL,document TEXT NOT NULL,original BLOB)")
        conn.execute("CREATE TABLE IF NOT EXISTS hc_workbook_history(id INTEGER PRIMARY KEY,revision INTEGER,document TEXT,original BLOB,created_at TEXT DEFAULT CURRENT_TIMESTAMP)")
        conn.commit();return conn
    @lru_cache(maxsize=2)
    def legacy_sheets(raw):
        return import_excel(raw)['sheets']
    def state(conn):
        row=conn.execute("SELECT * FROM hc_workbook WHERE id=1").fetchone()
        if row:
            document=json.loads(row["document"])
            if 'sheets' not in document:document['sheets']=copy.deepcopy(legacy_sheets(bytes(row['original']))) if row['original'] else sheet_defaults()
            return document,row["revision"],row["original"]
        return {"opening":0,"limit":5700000,"rows":[],"offsets":{"Payment":[],"Receipt":[]}},0,None
    def csrf():
        token=request.headers.get("X-HC-CSRF","")
        return secrets.compare_digest(str(session.get("hc_csrf","")),token) and bool(token)
    def save(conn,document,expected,original=None,replace_original=False):
        conn.execute("BEGIN IMMEDIATE")
        old=conn.execute("SELECT * FROM hc_workbook WHERE id=1").fetchone()
        revision=old["revision"] if old else 0
        if expected!=revision:
            conn.rollback();return False
        if old:
            conn.execute("INSERT INTO hc_workbook_history(revision,document,original) VALUES(?,?,?)",(revision,old["document"],old["original"]))
            conn.execute("DELETE FROM hc_workbook_history WHERE id NOT IN (SELECT id FROM hc_workbook_history ORDER BY id DESC LIMIT 10)")
        if old and not replace_original:original=old["original"]
        conn.execute("INSERT OR REPLACE INTO hc_workbook VALUES(1,?,?,?)",(revision+1,json.dumps(document),original));conn.commit();return True
    @bp.get("/hc-ledger")
    @require_permission("cash_view")
    def page(): return send_from_directory(ROOT,"hc_ledger.html")
    @bp.get("/hc-ledger-assets/<path:filename>")
    @require_permission("cash_view")
    def asset(filename):
        if filename not in ("hc_ledger.js","hc_ledger.css","hc_sheets.js"):return "",404
        return send_from_directory(ROOT,filename)
    @bp.get('/hc-sheets')
    @require_permission('cash_view')
    def linked_page():return send_from_directory(ROOT,'hc_sheets.html')
    @bp.get('/api/hc-workbook/sheet')
    @require_permission('cash_view')
    def linked_get():
        with closing(connect()) as conn:document,revision,_=state(conn)
        if not session.get('hc_csrf'):session['hc_csrf']=secrets.token_urlsafe(32)
        try:result=sheet_view(document,request.args.get('name','Deno HC'),calculate(document),number)
        except ValueError as exc:return jsonify(error=str(exc)),400
        return jsonify(**result,revision=revision,csrf=session['hc_csrf'])
    @bp.get("/api/hc-workbook")
    @require_permission("cash_view")
    def get():
        with closing(connect()) as conn:document,revision,_=state(conn)
        if not session.get("hc_csrf"):session["hc_csrf"]=secrets.token_urlsafe(32)
        return jsonify({**calculate(document),"opening":document["opening"],"limit":document["limit"],"offsets":ranked_offsets(document),"revision":revision,"template":template,"csrf":session["hc_csrf"]})
    @bp.post("/api/hc-workbook/import")
    @require_permission("cash_import")
    def import_file():
        if not csrf():return jsonify(error="Reload this page before saving"),403
        upload=request.files.get("file")
        if not upload or not upload.filename.lower().endswith(".xlsx"):return jsonify(error="Select the HC Ledger .xlsx file"),400
        raw=upload.stream.read(15*1024*1024+1)
        if len(raw)>15*1024*1024:return jsonify(error="Maximum upload size is 15 MB"),413
        try:
            document=import_excel(raw);revision=int(request.form.get("revision","-1"))
        except (ValueError,TypeError) as exc:return jsonify(error=str(exc)),400
        with closing(connect()) as conn:
            ok=save(conn,document,revision,raw,True)
        if not ok:return jsonify(error="Another user changed this ledger. Reload and try again."),409
        audit("Import","HC Workbook",f"{len(document['rows'])} rows imported")
        return jsonify(ok=True,revision=revision+1)
    @bp.get("/api/hc-workbook/history")
    @require_permission("cash_view")
    def history():
        with closing(connect()) as conn:
            entries=conn.execute("SELECT revision,created_at FROM hc_workbook_history ORDER BY id DESC").fetchall()
        return jsonify(entries=[dict(r) for r in entries])
    @bp.patch("/api/hc-workbook")
    @require_permission("cash_import")
    def update():
        if not csrf():return jsonify(error="Reload this page before saving"),403
        payload=request.get_json(silent=True) or {}
        try:
            with closing(connect()) as conn:
                document,revision,_=state(conn)
                action=payload.get("action");replacement=None;replace_original=False
                if action=='restore':
                    record=conn.execute('SELECT document,original FROM hc_workbook_history WHERE revision=? ORDER BY id DESC LIMIT 1',(int(payload['restore_revision']),)).fetchone()
                    if not record:raise ValueError('This history revision is no longer available')
                    document=json.loads(record['document']);replacement=record['original'];replace_original=True
                elif action=='sheet':
                    mutate_sheet(document,payload,number,date_text)
                elif action=='delete':
                    index=int(payload['row'])
                    if not 0<index<len(document['rows']):raise ValueError('The opening row cannot be deleted')
                    document['rows'][index]=['']*20
                elif action=="settings":
                    document["opening"]=number(payload["opening"]);document["limit"]=number(payload["limit"])
                    if document["limit"]<0:raise ValueError("Approved limit cannot be negative")
                elif action=="cell":
                    index=int(payload["row"]);col=int(payload["col"])
                    if not 0<=index<len(document["rows"]) or col not in INPUTS:raise ValueError("This cell is calculated or does not exist")
                    values=document["rows"][index].copy();values[col]=payload.get("value","")
                    candidate=sanitize_row(values)
                    if col in (3,4) and candidate[4] and candidate[4] not in document["offsets"].get(candidate[3],[]):raise ValueError("Select an offset account from the chosen Payment or Receipt list")
                    document["rows"][index]=candidate
                    if index==0 and (number(document["rows"][0][5]) or number(document["rows"][0][6])):raise ValueError("The first row is reserved for the opening balance")
                elif action=="append":
                    if len(document["rows"])>=9995:raise ValueError("Ledger has reached 9,995 entries")
                    row=sanitize_row(payload["values"])
                    if not document["rows"] and (number(row[5]) or number(row[6])):raise ValueError("Add an opening-balance row first")
                    if row[4] and row[4] not in document["offsets"].get(row[3],[]):raise ValueError("Select an offset account from the chosen Payment or Receipt list")
                    if (row[3]=="Payment" and number(row[5])) or (row[3]=="Receipt" and number(row[6])):raise ValueError("Enter the amount on the selected Payment or Receipt side")
                    document["rows"].append(row)
                else:raise ValueError("Unknown save action")
                ok=save(conn,document,payload.get("revision"),replacement,replace_original)
            if not ok:return jsonify(error="Another user changed this ledger. Reload before saving."),409
        except (ValueError,KeyError,TypeError,IndexError) as exc:return jsonify(error=str(exc)),400
        audit("Update","HC Workbook",str(action));return jsonify(ok=True,revision=revision+1)
    @bp.get("/api/hc-workbook/export")
    @require_permission("export_data")
    def export():
        with closing(connect()) as conn:document,_,original=state(conn)
        return send_file(export_excel(document,original,template),as_attachment=True,download_name="HC_Ledger_L4_Online.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    app.register_blueprint(bp)
